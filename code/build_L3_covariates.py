#!/usr/bin/env python3
"""Build UGR L3 inputs from the validated imaging handoff.

Run through build_L3_inputs.sh from this repository's code directory. The script
intersects the linked paper cohort with the validated session-01 UGR handoff,
preserves its selected one- or two-run input, and keeps subject order identical
across the selected-input manifest, generic FSL filepath list, frozen run
metrics, and the covariate workbook.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path


PROJECT = "UG / UGR"
TASK = "ugr"
EXPECTED_RUNS = ("1", "2")
EXPECTED_CONTRASTS = tuple(range(1, 18))
HANDOFF_SELECTION = "selected_coverage95_preserved_decisions"
SELECTION_POLICY = "validated_coverage95_handoff"
SUBJECT_RE = re.compile(r"^[0-9]+$")
HANDOFF_FIELDS = (
    "subject", "session", "analysis", "input_runs", "level", "contrast",
    "cope", "varcope", "mask", "dof", "feat_dir", "selection",
)
REQUIRED_ANALYSIS = (
    "rf1_id", "sub_age", "gender", "ecog_score", "mspss_sum", "adi_natrank",
    "pm25_mean_2013_2023",
)


def die(message: str) -> None:
    raise SystemExit(f"ERROR: {message}")


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: object) -> float | None:
    try:
        result = float(text(value))
    except ValueError:
        return None
    return result if math.isfinite(result) else None


def center(values: list[float]) -> list[float]:
    mean = sum(values) / len(values)
    return [value - mean for value in values]


def load_analysis(path: Path) -> tuple[list[str], dict[str, dict[str, object]], dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        die("Python package 'openpyxl' is required")
    if not path.is_file():
        die(f"analysis workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(rows)]
    except StopIteration:
        die(f"analysis workbook is empty: {path}")
    missing = [column for column in REQUIRED_ANALYSIS if column not in headers]
    if missing:
        die(f"analysis workbook is missing columns: {', '.join(missing)}")
    indices = {name: headers.index(name) for name in headers}
    order: list[str] = []
    data: dict[str, dict[str, object]] = {}
    problems: dict[str, str] = {}
    for excel_row, values in enumerate(rows, start=2):
        subject = text(values[indices["rf1_id"]]).removeprefix("sub-")
        if not SUBJECT_RE.fullmatch(subject):
            die(f"invalid rf1_id on Excel row {excel_row}: {subject!r}")
        if subject in data or subject in problems:
            die(f"duplicate rf1_id in analysis workbook: {subject}")
        row = {name: values[index] for name, index in indices.items()}
        missing_values = [
            column for column in REQUIRED_ANALYSIS
            if column != "rf1_id" and not text(row[column])
        ]
        invalid_numbers = [
            column for column in (
                "sub_age", "ecog_score", "mspss_sum", "adi_natrank",
                "pm25_mean_2013_2023",
            )
            if text(row[column]) and number(row[column]) is None
        ]
        if missing_values or invalid_numbers:
            details = [*(f"missing_{x}" for x in missing_values), *(f"invalid_{x}" for x in invalid_numbers)]
            problems[subject] = ";".join(details)
        else:
            order.append(subject)
            data[subject] = row
    workbook.close()
    if not order:
        die("analysis workbook has no participants with complete required covariates")
    return order, data, problems


def load_run_decisions(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    if not path.is_file():
        die(f"validated handoff run-decision table not found: {path}")
    selected: dict[tuple[str, str], dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {
            "subject", "session", "task", "run", "selection_status",
            "selection_reason", "coverage_pct", "coverage_decision",
            "flag_or_metric_tsnr", "flag_or_metric_fd_mean",
        }
        missing = sorted(required - set(reader.fieldnames or ()))
        if missing:
            die(f"run-decision table is missing columns: {', '.join(missing)}")
        for row in reader:
            subject = text(row.get("subject")).removeprefix("sub-")
            session = text(row.get("session")).removeprefix("ses-").zfill(2)
            task = text(row.get("task")).lower()
            run = text(row.get("run")).removeprefix("run-")
            if session != "01" or task != TASK:
                continue
            key = (subject, run)
            if key in selected:
                die(f"duplicate run-decision row for sub-{subject}, {TASK}, run-{run}")
            selected[key] = row
    return selected


def require_readable_file(value: str, subject: str, contrast: int, field: str) -> None:
    path = Path(value)
    try:
        if not path.is_file() or path.stat().st_size == 0:
            die(
                f"handoff {field} is missing or empty for sub-{subject}, "
                f"contrast {contrast}: {path}. This is an access/processing error, "
                "not a participant exclusion"
            )
    except OSError as exc:
        die(
            f"cannot access handoff {field} for sub-{subject}, contrast {contrast}: "
            f"{path} ({exc}). This is not a participant exclusion"
        )


def load_selected_inputs(path: Path) -> tuple[list[str], dict[str, list[dict[str, str]]]]:
    if not path.is_file():
        die(f"validated selected-input table not found: {path}")
    by_subject: dict[str, list[dict[str, str]]] = {}
    subject_order: list[str] = []
    seen_rows: set[tuple[str, int]] = set()
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        missing = sorted(set(HANDOFF_FIELDS) - set(reader.fieldnames or ()))
        if missing:
            die(f"selected-input table is missing columns: {', '.join(missing)}")
        for row in reader:
            subject = text(row.get("subject")).removeprefix("sub-")
            if not SUBJECT_RE.fullmatch(subject):
                die(f"invalid handoff subject: {subject!r}")
            if text(row.get("session")).removeprefix("ses-").zfill(2) != "01":
                die(f"unexpected session in UGR handoff for sub-{subject}")
            if text(row.get("analysis")).lower() != TASK:
                die(f"wrong analysis in UGR handoff for sub-{subject}")
            if text(row.get("selection")) != HANDOFF_SELECTION:
                die(f"unvalidated selection label in handoff for sub-{subject}")
            runs = text(row.get("input_runs"))
            level = text(row.get("level"))
            if runs not in {"1", "2", "1+2"} or level != ("L2" if runs == "1+2" else "L1"):
                die(f"unexpected UGR run subset/level for sub-{subject}: {runs}/{level}")
            try:
                contrast = int(text(row.get("contrast")))
            except ValueError:
                die(f"invalid contrast in handoff for sub-{subject}: {row.get('contrast')!r}")
            if contrast not in EXPECTED_CONTRASTS or (subject, contrast) in seen_rows:
                die(f"unexpected or duplicate contrast {contrast} for sub-{subject}")
            seen_rows.add((subject, contrast))
            for field in ("cope", "varcope", "mask", "dof"):
                value = text(row.get(field))
                if not value or "COPENUM" in value:
                    die(f"handoff {field} is not an exact path for sub-{subject}, contrast {contrast}")
                require_readable_file(value, subject, contrast, field)
            if subject not in by_subject:
                by_subject[subject] = []
                subject_order.append(subject)
            by_subject[subject].append(row)
    if not by_subject:
        die("selected-input table contains no UGR rows")
    for subject, rows in by_subject.items():
        rows.sort(key=lambda row: int(text(row["contrast"])))
        found = tuple(int(text(row["contrast"])) for row in rows)
        run_subsets = {text(row["input_runs"]) for row in rows}
        levels = {text(row["level"]) for row in rows}
        if found != EXPECTED_CONTRASTS or len(run_subsets) != 1 or len(levels) != 1:
            die(f"sub-{subject} does not have one internally consistent row for every UGR contrast")
    return subject_order, by_subject


def write_inputs(args: argparse.Namespace) -> None:
    order, _, covariate_problems = load_analysis(args.analysis)
    if covariate_problems:
        die(
            "the post-PM2.5 analysis workbook contains incomplete analysis variables; "
            "rerun the upstream cohort pipeline"
        )
    handoff_order, handoff = load_selected_inputs(args.selected_inputs)
    decisions = load_run_decisions(args.run_decisions)
    analysis_subjects = set(order)
    missing_from_handoff = [subject for subject in order if subject not in handoff]
    if missing_from_handoff:
        die(
            "post-PM2.5 participant(s) are absent from the validated UGR imaging pool: "
            + ", ".join(missing_from_handoff)
        )
    final_subjects = [subject for subject in handoff_order if subject in analysis_subjects]
    exclusions: list[tuple[str, str, str]] = []
    if not final_subjects:
        die("the linked cohort and validated UGR handoff do not overlap")

    final_qc: list[tuple[object, ...]] = []
    for subject in final_subjects:
        representative = handoff[subject][0]
        runs = text(representative["input_runs"])
        selected_rows: list[dict[str, str]] = []
        for run in runs.split("+"):
            row = decisions.get((subject, run))
            if row is None or text(row.get("selection_status")) != "eligible":
                die(f"selected handoff participant sub-{subject} lacks an eligible run-{run} decision")
            selected_rows.append(row)
        metrics: list[tuple[float, float, float]] = []
        for row in selected_rows:
            values = (
                number(row.get("flag_or_metric_tsnr")),
                number(row.get("flag_or_metric_fd_mean")),
                number(row.get("coverage_pct")),
            )
            if any(value is None for value in values):
                die(f"selected handoff participant sub-{subject} lacks frozen run metrics")
            metrics.append(values)
        final_qc.append((
            subject,
            sum(value[0] for value in metrics) / len(metrics),
            sum(value[1] for value in metrics) / len(metrics),
            sum(value[2] for value in metrics) / len(metrics),
            runs,
            text(representative["level"]),
            SELECTION_POLICY,
        ))

    args.subjects_output.write_text("".join(f"{x}\n" for x in final_subjects), encoding="utf-8")
    with args.selected_output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=HANDOFF_FIELDS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for subject in final_subjects:
            writer.writerows(handoff[subject])
    with args.final_qc_output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(("subject", "tsnr", "fd_mean", "brain_coverage_pct", "runs", "level", "selection_policy"))
        writer.writerows(final_qc)
    with args.exclusions_output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(("subject", "stage", "reason"))
        writer.writerows(exclusions)
    print(f"{PROJECT} analysis participants considered: {len(order) + len(covariate_problems)}")
    print(f"Final UGR L3 participants carried forward from PM2.5: {len(final_subjects)}")
    print(f"Participants excluded before L3: {len({row[0] for row in exclusions})}")


def read_subjects(path: Path) -> list[str]:
    values = [line.strip().removeprefix("sub-") for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if not values or len(values) != len(set(values)) or any(not SUBJECT_RE.fullmatch(x) for x in values):
        die(f"subject list is empty, duplicated, or invalid: {path}")
    return values


def read_final_qc(path: Path) -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            subject = text(row.get("subject"))
            if subject in result:
                die(f"duplicate final-QC row for sub-{subject}")
            result[subject] = row
    return result


def software_versions(bids_root: Path, subjects: list[str], qc: dict[str, dict[str, str]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for subject in subjects:
        runs = [value for value in text(qc[subject].get("runs")).split("+") if value]
        found: set[str] = set()
        for run in runs:
            sidecar = bids_root / f"sub-{subject}" / "ses-01" / "func" / f"sub-{subject}_ses-01_task-{TASK}_run-{run}_echo-2_part-mag_bold.json"
            if not sidecar.is_file():
                die(f"BIDS sidecar not found: {sidecar}")
            metadata = json.loads(sidecar.read_text(encoding="utf-8"))
            version = text(metadata.get("SoftwareVersions"))
            if not version:
                die(f"SoftwareVersions is missing from {sidecar}")
            found.add(version)
        if len(found) != 1:
            die(f"SoftwareVersions disagree across selected runs for sub-{subject}: {sorted(found)}")
        result[subject] = next(iter(found))
    return result


def normalized_gender(value: object) -> str:
    label = text(value).casefold().replace("_", "-").replace(" ", "-")
    mapping = {"0": "male", "male": "male", "1": "female", "female": "female", "2": "nonbinary", "nonbinary": "nonbinary", "non-binary": "nonbinary"}
    if label not in mapping:
        die(f"unsupported gender code/label: {text(value)!r}")
    return mapping[label]


def software_design(versions: dict[str, str], subjects: list[str]) -> tuple[str, dict[str, str], dict[str, list[float]]]:
    counts = Counter(versions[x] for x in subjects)
    reference = sorted(counts, key=lambda x: (-counts[x], x.casefold(), x))[0]
    columns: dict[str, str] = {}
    used: set[str] = set()
    for version in sorted(x for x in counts if x != reference):
        slug = re.sub(r"[^a-z0-9]+", "_", version.casefold()).strip("_") or "unnamed"
        base = f"software_version_{slug[:48]}_mc"
        name, suffix = base, 2
        while name in used:
            name, suffix = f"{base}_{suffix}", suffix + 1
        used.add(name)
        columns[version] = name
    values = {columns[version]: center([1.0 if versions[x] == version else 0.0 for x in subjects]) for version in columns}
    return reference, columns, values


def write_workbook(args: argparse.Namespace) -> None:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        die("Python package 'openpyxl' is required")
    subjects = read_subjects(args.subjects)
    _, analysis, problems = load_analysis(args.analysis)
    qc = read_final_qc(args.final_qc)
    for subject in subjects:
        if subject not in analysis:
            die(f"final participant sub-{subject} lacks complete analysis covariates: {problems.get(subject, 'not_found')}")
        if subject not in qc:
            die(f"final participant sub-{subject} lacks final-QC covariates")
    if set(qc) != set(subjects):
        die("final-QC subject membership does not exactly match the subject list")
    versions = software_versions(args.bids_root, subjects, qc)
    raw = [analysis[x] for x in subjects]
    ecog = center([number(x["ecog_score"]) for x in raw])
    adi = center([number(x["adi_natrank"]) for x in raw])
    mspss = center([number(x["mspss_sum"]) for x in raw])
    pm25 = center([number(x["pm25_mean_2013_2023"]) for x in raw])
    age = center([number(x["sub_age"]) for x in raw])
    tsnr = center([number(qc[x]["tsnr"]) for x in subjects])
    fd_mean = center([number(qc[x]["fd_mean"]) for x in subjects])
    genders = [normalized_gender(x["gender"]) for x in raw]
    male = center([1.0 if x == "male" else 0.0 for x in genders])
    female = center([1.0 if x == "female" else 0.0 for x in genders])
    reference, software_columns, software_values = software_design(versions, subjects)
    software_headers = [software_columns[x] for x in software_columns]
    headers = [
        "sub_id", "software_version", *software_headers, "tsnr", "fd_mean",
        "ecog_score_mc", "adi_natrank_mc", "mspss_sum_mc",
        "pm25_mean_2013_2023_mc", "sub_age_mc", "male", "female",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "covariates"
    sheet.append(headers)
    for index, subject in enumerate(subjects):
        sheet.append([
            subject, versions[subject],
            *(software_values[column][index] for column in software_headers),
            tsnr[index], fd_mean[index], ecog[index], adi[index], mspss[index],
            pm25[index], age[index], male[index], female[index],
        ])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(headers) + 1):
        letter = sheet.cell(1, column).column_letter
        sheet.column_dimensions[letter].width = 30 if headers[column - 1] == "software_version" else 22
        for cell in sheet[letter][1:]:
            cell.number_format = "@" if headers[column - 1] in {"sub_id", "software_version"} else "0.000000"
    metadata = workbook.create_sheet("coding")
    metadata.append(["item", "value", "detail"])
    metadata.append(["project", PROJECT, "Session 01 UGR activation"])
    metadata.append(["imaging_selection", SELECTION_POLICY, "Validated exact input rows intersected with the linked paper cohort"])
    metadata.append([
        "mspss", "mspss_sum_mc",
        "MSPSS total score, grand-mean centered within the final imaging sample",
    ])
    metadata.append(["pm25", "pm25_mean_2013_2023_mc", "2013-2023 annual PM2.5 mean, grand-mean centered within the final imaging sample"])
    metadata.append(["gender", "male and female", "Grand-mean-centered indicators; non-binary is 0 in both before centering"])
    metadata.append(["software_reference", reference, "Most frequent SoftwareVersions level (alphabetical tie-break)"])
    counts = Counter(versions.values())
    for version in sorted(counts):
        metadata.append(["software_level", version, f"n={counts[version]}; EV={software_columns.get(version, 'reference')}"])
    for cell in metadata[1]:
        cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
    metadata.freeze_panes = "A2"
    metadata.column_dimensions["A"].width = 24
    metadata.column_dimensions["B"].width = 34
    metadata.column_dimensions["C"].width = 80
    workbook.save(args.output)
    checked = load_workbook(args.output, read_only=True, data_only=False)
    if checked["covariates"].max_row != len(subjects) + 1:
        die("saved covariate workbook failed row-count verification")
    checked.close()
    print(f"Wrote {len(subjects)} aligned centered-covariate rows to {args.output}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    inputs = commands.add_parser("inputs")
    inputs.add_argument("--analysis", type=Path, required=True)
    inputs.add_argument("--selected-inputs", type=Path, required=True)
    inputs.add_argument("--run-decisions", type=Path, required=True)
    inputs.add_argument("--subjects-output", type=Path, required=True)
    inputs.add_argument("--selected-output", type=Path, required=True)
    inputs.add_argument("--final-qc-output", type=Path, required=True)
    inputs.add_argument("--exclusions-output", type=Path, required=True)
    inputs.set_defaults(function=write_inputs)
    workbook = commands.add_parser("workbook")
    workbook.add_argument("--analysis", type=Path, required=True)
    workbook.add_argument("--subjects", type=Path, required=True)
    workbook.add_argument("--final-qc", type=Path, required=True)
    workbook.add_argument("--bids-root", type=Path, required=True)
    workbook.add_argument("--output", type=Path, required=True)
    workbook.set_defaults(function=write_workbook)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.function(args)


if __name__ == "__main__":
    main()
